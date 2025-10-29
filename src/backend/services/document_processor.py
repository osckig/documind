"""
Document processing service for extracting text from various file formats
"""
import os
from typing import List, Dict, Any, Optional, Tuple
import logging
from pathlib import Path
import pypdf
import docx
from PIL import Image
import pytesseract
import pdfplumber
from langchain.text_splitter import RecursiveCharacterTextSplitter
from config.settings import settings

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Service for processing documents"""

    def __init__(self):
        """Initialize document processor"""
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", ".", "!", "?", ",", " ", ""]
        )

        # Set Tesseract path if configured
        if settings.tesseract_path:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_path

    async def process_document(
        self,
        file_path: str,
        filename: str
    ) -> Tuple[str, Dict[str, Any]]:
        """Process document and extract text"""
        try:
            file_ext = Path(filename).suffix.lower()

            if file_ext == ".pdf":
                return await self.process_pdf(file_path)
            elif file_ext == ".docx":
                return await self.process_docx(file_path)
            elif file_ext == ".txt":
                return await self.process_txt(file_path)
            elif file_ext in [".png", ".jpg", ".jpeg"]:
                return await self.process_image(file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_ext}")

        except Exception as e:
            logger.error(f"Error processing document {filename}: {e}")
            raise

    async def process_pdf(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from PDF"""
        try:
            full_text = ""
            metadata = {"page_count": 0, "pages": []}

            with pdfplumber.open(file_path) as pdf:
                metadata["page_count"] = len(pdf.pages)

                for i, page in enumerate(pdf.pages):
                    # Extract text
                    text = page.extract_text()

                    if text:
                        full_text += f"\n\n--- Page {i + 1} ---\n\n{text}"
                        metadata["pages"].append({
                            "page_number": i + 1,
                            "char_count": len(text)
                        })

            metadata["total_chars"] = len(full_text)
            metadata["word_count"] = len(full_text.split())

            logger.info(f"Extracted {metadata['word_count']} words from PDF")
            return full_text.strip(), metadata

        except Exception as e:
            logger.error(f"Error processing PDF: {e}")
            raise

    async def process_docx(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from DOCX"""
        try:
            doc = docx.Document(file_path)
            full_text = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])

            metadata = {
                "paragraph_count": len(doc.paragraphs),
                "word_count": len(full_text.split()),
                "char_count": len(full_text)
            }

            logger.info(f"Extracted {metadata['word_count']} words from DOCX")
            return full_text, metadata

        except Exception as e:
            logger.error(f"Error processing DOCX: {e}")
            raise

    async def process_txt(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from TXT"""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                full_text = f.read()

            metadata = {
                "word_count": len(full_text.split()),
                "char_count": len(full_text),
                "line_count": len(full_text.split("\n"))
            }

            logger.info(f"Extracted {metadata['word_count']} words from TXT")
            return full_text, metadata

        except Exception as e:
            logger.error(f"Error processing TXT: {e}")
            raise

    async def process_image(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        """Extract text from image using OCR"""
        try:
            image = Image.open(file_path)

            # Perform OCR
            text = pytesseract.image_to_string(
                image,
                lang=settings.ocr_language
            )

            metadata = {
                "image_size": image.size,
                "image_mode": image.mode,
                "word_count": len(text.split()),
                "char_count": len(text)
            }

            logger.info(f"Extracted {metadata['word_count']} words from image using OCR")
            return text, metadata

        except Exception as e:
            logger.error(f"Error processing image: {e}")
            raise

    async def chunk_text(
        self,
        text: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> List[Dict[str, Any]]:
        """Split text into chunks"""
        try:
            # Split text into chunks
            chunks = self.text_splitter.split_text(text)

            # Create chunk objects with metadata
            chunk_objects = []
            for i, chunk in enumerate(chunks):
                chunk_obj = {
                    "content": chunk,
                    "chunk_index": i,
                    "char_count": len(chunk),
                    "word_count": len(chunk.split())
                }

                if metadata:
                    chunk_obj["metadata"] = metadata

                chunk_objects.append(chunk_obj)

            logger.info(f"Created {len(chunk_objects)} chunks from text")
            return chunk_objects

        except Exception as e:
            logger.error(f"Error chunking text: {e}")
            raise

    async def extract_keywords(self, text: str, top_k: int = 10) -> List[str]:
        """Extract keywords from text (simple implementation)"""
        try:
            # Simple keyword extraction based on frequency
            # In production, use more sophisticated methods like TF-IDF or RAKE
            words = text.lower().split()

            # Remove common stop words
            stop_words = set([
                "the", "a", "an", "and", "or", "but", "in", "on", "at", "to",
                "for", "of", "with", "by", "from", "as", "is", "was", "are",
                "were", "been", "be", "have", "has", "had", "do", "does", "did"
            ])

            # Count word frequency
            word_freq = {}
            for word in words:
                if word not in stop_words and len(word) > 3:
                    word_freq[word] = word_freq.get(word, 0) + 1

            # Get top k keywords
            keywords = sorted(word_freq.items(), key=lambda x: x[1], reverse=True)[:top_k]
            return [word for word, freq in keywords]

        except Exception as e:
            logger.error(f"Error extracting keywords: {e}")
            return []

    def get_preview(self, text: str, max_length: int = 200) -> str:
        """Get preview of text"""
        if len(text) <= max_length:
            return text
        return text[:max_length] + "..."
